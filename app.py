"""Streamlit UI for authenticated, isolated data analysis."""
import hashlib
import io
import logging
import os
import zipfile
import csv
import json
import time
import uuid

import pandas as pd
import streamlit as st
from dotenv import load_dotenv
from src.data_ops import left_join_tables

load_dotenv()
logger = logging.getLogger(__name__)
st.set_page_config(page_title="AI Data Analyst", page_icon="A", layout="wide")
USER_DATA_KEYS = ("df", "original_df", "dataframe_blob", "upload_digest",
                  "dataset_versions", "dataset_revision", "analysis_artifacts",
                  "pending_prompt", "dataset_dictionaries", "related_tables")


def _parse_uploaded_table(raw, filename):
    """Parse a bounded CSV/XLSX and reject malformed or ambiguous headers."""
    if filename.lower().endswith(".csv"):
        reader = csv.reader(io.StringIO(raw.decode("utf-8-sig")))
        header = next(reader, None)
        if not header:
            raise ValueError("The CSV is empty or has no header row.")
        if len(header) > 200:
            raise ValueError("Dataset exceeds the 200-column limit.")
        if len(header) != len(set(header)):
            raise ValueError("CSV contains duplicate column names.")
        row_count = sum(1 for _ in reader)
        if row_count > 500_000:
            raise ValueError("Dataset exceeds the 500,000-row limit.")
        frame = pd.read_csv(io.BytesIO(raw), nrows=500_001)
    else:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            members = archive.infolist()
            if (len(members) > 10_000 or sum(item.file_size for item in members) > 128 * 1024 * 1024
                    or any(item.file_size > 64 * 1024 * 1024
                           or (item.compress_size == 0 and item.file_size > 0)
                           or (item.compress_size and item.file_size / item.compress_size > 200)
                           for item in members)):
                raise ValueError("Workbook archive exceeds decompression limits.")
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        if sheet.max_column > 200 or sheet.max_row > 500_001:
            workbook.close()
            raise ValueError("Dataset exceeds the 500,000-row or 200-column limit.")
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None) or []
        workbook.close()
        cleaned = [item for item in header if item is not None]
        if len(cleaned) != len(set(cleaned)):
            raise ValueError("Excel sheet contains duplicate column names.")
        frame = pd.read_excel(io.BytesIO(raw), nrows=500_001)
    if frame.empty or len(frame.columns) == 0:
        raise ValueError("The dataset has no usable rows or columns.")
    if len(frame) > 500_000 or len(frame.columns) > 200:
        raise ValueError("Dataset exceeds the 500,000-row or 200-column limit.")
    if not frame.columns.is_unique:
        raise ValueError("Dataset contains duplicate column names.")
    frame = frame.loc[:, ~frame.columns.astype(str).str.startswith("Unnamed:")]
    if frame.empty or len(frame.columns) == 0:
        raise ValueError("The dataset has no usable columns.")
    return frame


def _notebook_for_session(messages, artifacts):
    cells = [{"cell_type": "markdown", "metadata": {}, "source": [
        "# AI Data Analyst session\n",
        "Download `analysis_dataset.parquet` beside this notebook before running it.\n"]},
        {"cell_type": "code", "execution_count": None, "metadata": {}, "outputs": [],
         "source": ["import pandas as pd\n", "import numpy as np\n",
                    "import plotly.express as px\n", "import plotly.graph_objects as go\n",
                    "from plotly.subplots import make_subplots\n",
                    "from src.agents.toolkit import stats, ts, viz\n",
                    "df = pd.read_parquet('analysis_dataset.parquet')\n",
                    "_session_baseline = df.copy(deep=True)\n"]}]
    for index, message in enumerate(messages):
        artifact = artifacts.get(index)
        if not artifact:
            continue
        question = messages[index - 1].content if index and hasattr(messages[index - 1], "content") else "Analysis"
        cells.append({"cell_type": "markdown", "metadata": {},
                      "source": [f"## {question}\n", f"Dataset version: {artifact.get('data_version', 'unknown')}\n"]})
        if artifact.get("code") and artifact.get("intent") == "MODIFY":
            cells.append({"cell_type": "markdown", "metadata": {}, "source": [
                "Dataset modification already present in the downloaded current dataset; "
                "the modification code is preserved in the app's analysis details but not replayed here.\n"]})
        elif artifact.get("code"):
            cells.append({"cell_type": "code", "execution_count": None, "metadata": {}, "outputs": [],
                          "source": ["df = _session_baseline.copy(deep=True)\n",
                                     *[line + "\n" for line in artifact["code"].splitlines()] ]})
    return json.dumps({"cells": cells, "metadata": {"kernelspec": {
        "display_name": "Python 3", "language": "python", "name": "python3"}},
        "nbformat": 4, "nbformat_minor": 5}, ensure_ascii=False, indent=2)

if not os.getenv("GROQ_API_KEY"):
    st.error("GROQ_API_KEY is missing. Configure it before starting the app.")
    st.stop()

try:
    from langchain_core.messages import AIMessage, HumanMessage
    from src.agents.graph import app as agent_app
    from src.auth import (sign_in_with_email_and_password,
                          sign_up_with_email_and_password, validate_session,
                          revoke_session, consume_analysis_quota, AuthUnavailable,
                          QuotaExceeded)
except ImportError:
    logger.exception("Application module import failed")
    st.error("The application could not start because a required component is unavailable.")
    st.stop()


st.markdown("""<style>
:root {
  --canvas: #0b0f14;
  --panel: #121820;
  --panel-raised: #171f29;
  --line: rgba(173, 190, 210, .14);
  --text: #edf3f8;
  --muted: #9aa8b7;
  --accent: #27c2d5;
}
.stApp { background: radial-gradient(ellipse at 70% -20%, #15303a 0, var(--canvas) 46%); color: var(--text); }
[data-testid="stHeader"] { background: transparent; }
main .block-container { max-width: 1160px; padding: 2.5rem 2.5rem 5rem; }
h1, h2, h3 { color: var(--text) !important; letter-spacing: -.035em; }
h1 { font-size: clamp(2rem, 4vw, 3rem) !important; font-weight: 750 !important; }
[data-testid="stSidebar"] { background: #10161e; border-right: 1px solid var(--line); }
[data-testid="stSidebar"] > div:first-child { padding-top: 1.2rem; }
[data-testid="stSidebar"] h1, [data-testid="stSidebar"] h2,
[data-testid="stSidebar"] h3 { color: var(--text) !important; }
.workspace-brand { display:flex; align-items:center; gap:.7rem; margin: .15rem 0 1.6rem; }
.workspace-mark { display:grid; place-items:center; width:2.35rem; height:2.35rem; border-radius:.75rem;
  background:linear-gradient(135deg,#33d2df,#148ba7); color:#07151b; font-weight:900; font-size:1.2rem; }
.workspace-name { font-size:1rem; font-weight:700; color:var(--text); line-height:1.2; }
.workspace-kicker { font-size:.65rem; letter-spacing:.14em; color:var(--muted); margin-top:.2rem; }
.page-kicker { color:var(--accent); text-transform:uppercase; letter-spacing:.16em; font-size:.72rem; font-weight:700; margin:0 0 .35rem; }
.page-subtitle { color:var(--muted); font-size:1.03rem; margin-top:-.45rem; margin-bottom:1.7rem; }
.dataset-status { display:inline-flex; align-items:center; gap:.45rem; border:1px solid var(--line);
  background:rgba(18,24,32,.8); color:#b8c6d3; padding:.35rem .7rem; border-radius:999px;
  font-size:.75rem; margin:.2rem 0 1rem; }
.dataset-dot { width:.45rem; height:.45rem; border-radius:50%; background:#39d39a; }
.stChatMessage { background:rgba(23,31,41,.88); border:1px solid var(--line); border-radius:1rem;
  padding:1.05rem 1.15rem; margin:.8rem 0; }
[data-testid="stChatInput"] textarea { border:1px solid var(--line); border-radius:.85rem;
  background:#121820; color:var(--text); }
[data-testid="stChatInput"] textarea:focus { border-color:var(--accent); box-shadow:0 0 0 1px var(--accent); }
[data-testid="stFileUploader"] section { background:#121820; border:1px dashed #43515e; border-radius:.9rem; }
[data-testid="stExpander"] { background:rgba(18,24,32,.68); border:1px solid var(--line); border-radius:.8rem; }
[data-testid="stMetric"] { background:var(--panel-raised); border:1px solid var(--line); border-radius:.8rem; padding:.65rem .8rem; }
div.stButton > button, [data-testid="stFormSubmitButton"] button { border-radius:.7rem; font-weight:650; }
div.stButton > button:hover { border-color:var(--accent); color:var(--accent); }
@media (max-width: 700px) { main .block-container { padding:1.5rem 1rem 4rem; } }
</style>""", unsafe_allow_html=True)

if "user" not in st.session_state:
    _, auth_col, _ = st.columns([1, 1.15, 1])
    with auth_col:
        st.markdown('<div class="workspace-brand"><div class="workspace-mark">A</div><div><div class="workspace-name">AI Data Analyst</div><div class="workspace-kicker">PRIVATE DATA WORKSPACE</div></div></div>', unsafe_allow_html=True)
        st.markdown('<p class="page-kicker">Welcome</p>', unsafe_allow_html=True)
        st.title("Sign in to continue")
        st.markdown('<p class="page-subtitle">Your datasets and analysis history stay in your workspace.</p>', unsafe_allow_html=True)
        auth_mode = st.radio("Account", ["Login", "Sign Up"], horizontal=True, label_visibility="collapsed")
        with st.form("auth_form"):
            email = st.text_input("Email", max_chars=254)
            password = st.text_input("Password", type="password", max_chars=128,
                                     help="Use at least 12 characters.")
            invite_code = (st.text_input("Invitation code", type="password", max_chars=128)
                           if auth_mode == "Sign Up" and os.getenv("SIGNUP_INVITE_CODE") else None)
            submitted = st.form_submit_button("Continue", use_container_width=True, type="primary")
        if submitted:
            auth_result = (sign_in_with_email_and_password(email, password)
                           if auth_mode == "Login"
                           else sign_up_with_email_and_password(email, password, invite_code))
            if auth_result.get("session_token"):
                st.session_state["user"] = auth_result
                st.rerun()
            st.error(auth_result.get("error", {}).get("message", "Authentication failed."))
    st.stop()

try:
    validated_user = validate_session(st.session_state["user"])
except AuthUnavailable:
    st.error("Sign-in verification is temporarily unavailable. Please retry in a moment.")
    st.stop()
if not validated_user:
    st.session_state.pop("user", None)
    for key in USER_DATA_KEYS:
        st.session_state.pop(key, None)
    st.session_state["upload_widget_version"] = st.session_state.get("upload_widget_version", 0) + 1
    st.session_state["messages"] = [AIMessage(content="Upload a dataset and ask me anything about it.")]
    st.warning("Your session expired or could not be verified. Please sign in again.")
    st.rerun()
st.session_state["user"] = validated_user

frame = st.session_state.get("df")
st.markdown('<p class="page-kicker">AI-powered data exploration</p>', unsafe_allow_html=True)
st.title("Make sense of your data.")
st.markdown('<p class="page-subtitle">Ask a question in plain English. Get a clear answer, analysis, and charts.</p>', unsafe_allow_html=True)
if frame is not None:
    st.markdown(f'<div class="dataset-status"><span class="dataset-dot"></span> Dataset ready <span>·</span> {len(frame):,} rows <span>·</span> {len(frame.columns):,} columns</div>', unsafe_allow_html=True)
else:
    st.markdown('<div class="dataset-status"><span class="dataset-dot" style="background:#788797"></span> Upload a dataset to get started</div>', unsafe_allow_html=True)
if "messages" not in st.session_state:
    st.session_state["messages"] = [AIMessage(content="Upload a dataset and ask me anything about it.")]
if "analysis_artifacts" not in st.session_state:
    st.session_state["analysis_artifacts"] = {}

# Upload handling is deliberately below the auth gate.
with st.sidebar:
    st.markdown('<div class="workspace-brand"><div class="workspace-mark">A</div><div><div class="workspace-name">AI Data Analyst</div><div class="workspace-kicker">YOUR WORKSPACE</div></div></div>', unsafe_allow_html=True)
    st.markdown("#### Privacy")
    privacy_tier = st.selectbox("Dataset context shared with Groq", [
        "Tier 0 · Schema only", "Tier 1 · Local profile", "Tier 2 · Include sample rows"],
        index=0, help="Column names, types, shape, and your question are always sent to Groq. Tier 2 also shares sample rows and analysis output during planning, code generation, and summaries.")
    share_values = privacy_tier.startswith("Tier 2")
    st.markdown("#### Dataset")
    uploader_key = f"dataset_upload_{st.session_state.get('upload_widget_version', 0)}"
    uploaded_files = st.file_uploader("Upload CSV or Excel tables", type=["csv", "xlsx"],
                                      accept_multiple_files=True, key=uploader_key,
                                      help="Up to 5 tables, 20 MB combined; CSV or XLSX.")
    if uploaded_files:
        try:
            if len(uploaded_files) > 5:
                raise ValueError("Upload at most 5 related tables.")
            uploads = [(file.name, file.getvalue()) for file in uploaded_files]
            if sum(len(raw) for _, raw in uploads) > 20 * 1024 * 1024:
                raise ValueError("Combined upload size exceeds the 20 MB limit.")
            digest = hashlib.sha256("|".join(name + ":" + hashlib.sha256(raw).hexdigest()
                                               for name, raw in uploads).encode()).hexdigest()
            if st.session_state.get("upload_digest") != digest:
                tables = {}
                for index, (filename, raw) in enumerate(uploads):
                    if not raw:
                        raise ValueError(f"{filename} is empty.")
                    name = filename.rsplit(".", 1)[0][:80] or f"table_{index + 1}"
                    if name in tables:
                        name = f"{name}_{index + 1}"
                    table = _parse_uploaded_table(raw, filename)
                    encoded = io.BytesIO()
                    table.to_parquet(encoded, index=None, compression="zstd")
                    max_bytes = int(os.getenv("MAX_DATASET_PARQUET_BYTES", str(16 * 1024 * 1024)))
                    if encoded.tell() > max_bytes:
                        raise ValueError(f"{filename} converts to {encoded.tell() / 1024 / 1024:.1f} MB; per-table limit is {max_bytes / 1024 / 1024:.1f} MB.")
                    tables[name] = table
                frame = next(iter(tables.values()))
                encoded = io.BytesIO()
                frame.to_parquet(encoded, index=None, compression="zstd")
                st.session_state["related_tables"] = tables
                st.session_state["df"] = frame
                st.session_state["original_df"] = frame.copy(deep=True)
                st.session_state["dataframe_blob"] = encoded.getvalue()
                st.session_state["dataset_versions"] = []
                st.session_state["dataset_revision"] = 1
                st.session_state["upload_digest"] = digest
                st.session_state["analysis_artifacts"] = {}
                st.session_state["messages"] = [AIMessage(content="Your tables are ready. Ask a question or join related tables in the sidebar.")]
                st.success(f"Loaded {len(tables)} table(s); starting with {next(iter(tables))}.")
        except Exception as exc:
            logger.warning("Dataset upload rejected or failed (%s)", type(exc).__name__)
            st.error(str(exc) if isinstance(exc, ValueError) else "One of these files could not be read. Check that each is a valid CSV or Excel workbook.")

    tables = st.session_state.get("related_tables", {})
    if len(tables) > 1:
        st.markdown("#### Join tables")
        left_name = st.selectbox("Left table", list(tables), key="join_left_table")
        right_options = [name for name in tables if name != left_name]
        right_name = st.selectbox("Right table", right_options, key="join_right_table")
        left_cols, right_cols = list(tables[left_name].columns), list(tables[right_name].columns)
        common = [column for column in left_cols if column in right_cols]
        left_key = st.selectbox("Left join key", left_cols,
                                index=left_cols.index(common[0]) if common else 0)
        right_key = st.selectbox("Right join key", right_cols,
                                 index=right_cols.index(left_key) if left_key in right_cols else (right_cols.index(common[0]) if common else 0))
        if st.button("Create joined dataset", type="primary", use_container_width=True):
            try:
                max_bytes = int(os.getenv("MAX_DATASET_PARQUET_BYTES", str(16 * 1024 * 1024)))
                joined, parquet_blob = left_join_tables(
                    tables[left_name], tables[right_name], left_key, right_key,
                    max_parquet_bytes=max_bytes)
                versions = st.session_state.setdefault("dataset_versions", [])
                versions.append({"df": st.session_state["df"].copy(deep=True),
                                 "blob": st.session_state.get("dataframe_blob"),
                                 "revision": st.session_state.get("dataset_revision", 1)})
                st.session_state["dataset_versions"] = versions[-5:]
                st.session_state["df"] = joined
                st.session_state["dataframe_blob"] = parquet_blob
                st.session_state["dataset_revision"] += 1
                st.success(f"Joined {left_name} to {right_name} ({len(joined):,} rows).")
                st.rerun()
            except Exception as exc:
                logger.warning("Table join failed (%s)", type(exc).__name__)
                st.error(str(exc) if isinstance(exc, ValueError) else "Could not join these tables.")
    if "df" in st.session_state:
        frame = st.session_state["df"]
        dataset_key = st.session_state.get("upload_digest", "current")
        dictionaries = st.session_state.setdefault("dataset_dictionaries", {})
        dictionaries[dataset_key] = st.text_area(
            "Data dictionary / metric definitions", value=dictionaries.get(dataset_key, ""),
            key=f"dictionary_{dataset_key[:12]}", height=90, max_chars=3000,
            help="Describe column meanings, units, or business definitions. This text is shared with Groq.")
        for old_key in list(dictionaries)[:-10]:
            dictionaries.pop(old_key, None)
        st.markdown("#### Data health")
        metric_cols = st.columns(2)
        metric_cols[0].metric("Rows", f"{len(frame):,}")
        metric_cols[1].metric("Columns", f"{len(frame.columns):,}")
        if st.session_state.get("dataframe_blob"):
            st.download_button("Download current dataset (Parquet)",
                                st.session_state["dataframe_blob"],
                                file_name="analysis_dataset.parquet",
                                mime="application/vnd.apache.parquet",
                                use_container_width=True)
        if st.session_state.get("analysis_artifacts"):
            st.download_button("Export analysis notebook",
                                _notebook_for_session(st.session_state["messages"],
                                                      st.session_state["analysis_artifacts"]),
                                file_name="analysis_session.ipynb", mime="application/x-ipynb+json",
                                use_container_width=True,
                                help="Also download the current Parquet dataset and place it beside the notebook.")
        if st.button("Reset dataset to original", use_container_width=True):
            st.session_state["df"] = st.session_state["original_df"].copy(deep=True)
            parquet_buffer = io.BytesIO()
            st.session_state["df"].to_parquet(parquet_buffer, index=None, compression="zstd")
            st.session_state["dataframe_blob"] = parquet_buffer.getvalue()
            st.session_state["dataset_versions"] = []
            st.session_state["dataset_revision"] = 1
            st.rerun()
        versions = st.session_state.get("dataset_versions", [])
        if versions and st.button("Undo last dataset change", use_container_width=True):
            previous = versions.pop()
            st.session_state["dataset_versions"] = versions
            st.session_state["df"] = previous["df"]
            st.session_state["dataframe_blob"] = previous["blob"]
            st.session_state["dataset_revision"] = previous["revision"]
            st.rerun()
        st.caption(f"{int(frame.isnull().sum().sum()):,} missing values · {int(frame.duplicated().sum()):,} duplicate rows")
        with st.expander("Explore schema"):
            schema_view = pd.DataFrame({
                "Column": frame.columns.astype(str),
                "Type": frame.dtypes.astype(str).values,
                "Missing": frame.isnull().sum().values,
            })
            st.dataframe(schema_view, hide_index=True, use_container_width=True, height=280)
        with st.expander("Automatic overview and suggested questions"):
            constants = [str(col) for col in frame.columns if frame[col].nunique(dropna=True) <= 1]
            missing = frame.isna().mean().mul(100).sort_values(ascending=False)
            missing = missing[missing > 0].head(10)
            st.write(f"Rows: {len(frame):,}; columns: {len(frame.columns):,}; missing cells: {int(frame.isna().sum().sum()):,}; duplicate rows: {int(frame.duplicated().sum()):,}.")
            if constants:
                st.caption("Constant or nearly empty columns: " + ", ".join(constants[:20]))
            if not missing.empty:
                st.caption("Columns with missing values: " + ", ".join(
                    f"{column} ({value:.1f}%)" for column, value in missing.items()))
            numeric = frame.select_dtypes(include="number")
            if numeric.shape[1] >= 2:
                correlations = numeric.corr().abs().where(
                    ~pd.DataFrame([[i == j for j in range(len(numeric.columns))]
                                   for i in range(len(numeric.columns))],
                                  index=numeric.columns, columns=numeric.columns))
                pairs = correlations.stack().sort_values(ascending=False).head(3)
                if not pairs.empty:
                    st.caption("Strongest numeric relationships (absolute correlation): " + ", ".join(
                        f"{left} × {right}: {value:.2f}" for (left, right), value in pairs.items()))
            candidate = next((str(col) for col in frame.columns if frame[col].nunique(dropna=True) > 1), None)
            if candidate:
                st.write("Suggested questions")
                st.caption(f"How is {candidate} distributed? · Compare outcomes by {candidate} · Which columns have the most missing values?")

with st.sidebar:
    if st.button("Logout"):
        try:
            revoke_session(st.session_state.get("user"))
        except AuthUnavailable:
            logger.warning("Session could not be revoked because PostgreSQL is unavailable")
        for key in ("user", *USER_DATA_KEYS):
            st.session_state.pop(key, None)
        st.session_state["upload_widget_version"] = st.session_state.get("upload_widget_version", 0) + 1
        st.session_state["messages"] = [AIMessage(content="Upload a dataset and ask me anything about it.")]
        st.rerun()

if "pending_prompt" in st.session_state:
    prompt = st.session_state.pop("pending_prompt")
    if "df" not in st.session_state:
        st.session_state["messages"].append(
            AIMessage(content="Please upload a dataset first so I can analyze it."))
    else:
        frame = st.session_state["df"]
        data_version = st.session_state.get("dataset_revision", 1)
        schema = (f"Columns: {list(frame.columns)}\nTypes: {frame.dtypes.astype(str).to_dict()}\n"
                  f"Shape: {frame.shape}\nDataset version: {data_version}")
        if privacy_tier.startswith("Tier 1"):
            profile = []
            for col in frame.columns:
                series = frame[col]
                item = {"column": str(col), "dtype": str(series.dtype),
                        "null_pct": round(float(series.isna().mean() * 100), 2),
                        "cardinality": int(series.nunique(dropna=True))}
                if pd.api.types.is_numeric_dtype(series):
                    item.update({"min": str(series.min()), "median": str(series.median()), "max": str(series.max())})
                elif series.nunique(dropna=True) <= 20:
                    item["top_categories"] = series.value_counts(dropna=True).head(5).index.astype(str).tolist()
                profile.append(item)
            schema += f"\nLocal aggregate profile (no rows): {profile}"
        if share_values:
            schema += f"\nUser opted in to sharing dataset values. Sample rows: {frame.head(3).to_dict()}"
        dictionary = st.session_state.get("dataset_dictionaries", {}).get(st.session_state.get("upload_digest", "current"), "")
        if dictionary:
            schema += f"\nUser-provided data dictionary: {dictionary[:3000]}"
        state = {
            "messages": st.session_state["messages"], "dataset_schema": schema,
            "dataframe": frame, "dataframe_blob": st.session_state.get("dataframe_blob"), "plan": [], "code": "", "code_output": "",
            "explanation": "", "charts": [], "error": None, "retry_count": 0,
            "retry_history": [], "intent": "", "analysis_failed": False,
            "share_values": share_values,
        }
        try:
            if not consume_analysis_quota(st.session_state["user"]["user_id"]):
                raise QuotaExceeded("Daily analysis limit reached. Try again tomorrow.")
            run_id = uuid.uuid4().hex
            started = time.perf_counter()
            progress = st.status("Starting analysis…", expanded=True)
            result = dict(state)
            stages = {"query": "Classifying your request", "direct": "Preparing a response",
                      "planner": "Planning the analysis", "codegen": "Writing analysis code",
                      "executor": "Running in the isolated worker", "explainer": "Checking the result",
                      "failure": "Finalizing the failure report"}
            try:
                for update in agent_app.stream(state, stream_mode="updates"):
                    for node_name, node_result in update.items():
                        progress.update(label=stages.get(node_name, "Working"), state="running")
                        result.update(node_result)
                failed = bool(result.get("analysis_failed"))
                progress.update(label="Analysis failed" if failed else "Analysis finished",
                                state="error" if failed else "complete", expanded=False)
            except Exception:
                progress.update(label="Analysis could not be completed", state="error", expanded=False)
                raise
            logger.info(json.dumps({"event": "analysis_complete", "run_id": run_id,
                                    "intent": result.get("intent"),
                                    "attempts": result.get("retry_count", 0),
                                    "error_category": result.get("error_category"),
                                    "failed": bool(result.get("analysis_failed")),
                                    "elapsed_ms": round((time.perf_counter() - started) * 1000)}))
        except QuotaExceeded as exc:
            answer = str(exc)
        except Exception as exc:
            logger.error("Analysis request failed (%s)", type(exc).__name__)
            answer = "Analysis could not be completed. Please try again or check the service configuration."
        else:
            answer = result.get("explanation", "Analysis could not be completed.")
            if (not result.get("analysis_failed") and result.get("intent") == "MODIFY"
                    and result.get("dataframe") is not None):
                versions = st.session_state.get("dataset_versions", [])
                versions.append({"df": st.session_state["df"].copy(deep=True), "blob": st.session_state.get("dataframe_blob")})
                versions[-1]["revision"] = st.session_state.get("dataset_revision", 1)
                st.session_state["dataset_versions"] = versions[-5:]
                st.session_state["df"] = result["dataframe"]
                st.session_state["dataframe_blob"] = result.get("dataframe_blob")
                st.session_state["dataset_revision"] = st.session_state.get("dataset_revision", 1) + 1
            artifact = {
                "code_output": result.get("code_output", ""),
                "charts": result.get("charts", []),
                "plan": result.get("plan", []), "code": result.get("code", ""),
                "retry_history": result.get("retry_history", []),
                "intent": result.get("intent", ""), "data_version": data_version,
                "steps": [],
            }
            if (not result.get("analysis_failed")
                    and result.get("intent") in {"ANALYZE", "MODIFY"}
                    and (artifact["code_output"] or artifact["charts"] or artifact["code"])):
                st.session_state["analysis_artifacts"][len(st.session_state["messages"])] = artifact
                while len(st.session_state["analysis_artifacts"]) > 10:
                    st.session_state["analysis_artifacts"].pop(min(st.session_state["analysis_artifacts"]))
        st.session_state["messages"].append(AIMessage(content=answer))
    st.rerun()

for message_index, message in enumerate(st.session_state["messages"]):
    role = "user" if isinstance(message, HumanMessage) else "assistant"
    with st.chat_message(role):
        st.write(message.content)
    artifact = st.session_state["analysis_artifacts"].get(message_index)
    if artifact and artifact.get("intent") not in {"ANALYZE", "MODIFY"}:
        artifact = None
    if artifact:
        tab_names = ["Result", "Visualizations", "How it was computed"] if artifact["charts"] else ["Result", "How it was computed"]
        tabs = st.tabs(tab_names)
        with tabs[0]:
            if artifact["code_output"]:
                st.code(artifact["code_output"])
            elif not artifact["charts"]:
                st.caption("No additional computed output was returned.")
        if artifact["charts"]:
            with tabs[1]:
                for chart_index, chart in enumerate(artifact["charts"]):
                    chart_height = chart.layout.height
                    chart_height = 420 if chart_height is None or chart_height < 320 else min(chart_height, 620)
                    chart.update_layout(height=chart_height, autosize=True, margin={"l": 24, "r": 24, "t": 64, "b": 48})
                    st.plotly_chart(chart, use_container_width=True, theme="streamlit",
                                    config={"displaylogo": False, "responsive": True, "scrollZoom": True},
                                    key=f"analysis_chart_{message_index}_{chart_index}")
        detail_tab = tabs[-1]
        with detail_tab:
            st.caption(f"Dataset version: {artifact.get('data_version', 'unknown')} · Intent: {artifact.get('intent', 'analysis').lower()}")
            if artifact.get("plan"):
                st.markdown("**Plan**")
                for step in artifact["plan"]:
                    st.markdown(f"- {step}")
            if artifact["steps"]:
                with st.expander("Workflow messages"):
                    for item_type, item_content in artifact["steps"]:
                        st.caption(f"{item_type}: {item_content}")
            if artifact.get("code"):
                with st.expander("Generated code"):
                    st.code(artifact["code"], language="python")
                    st.download_button("Download analysis code", artifact["code"], file_name="analysis.py", key=f"code_download_{message_index}")
            if artifact.get("retry_history"):
                with st.expander("Execution attempts"):
                    for attempt in artifact["retry_history"]:
                        st.caption(f"Attempt {attempt.get('attempt')}: {attempt.get('category')} — {attempt.get('error') or 'completed'}")

if prompt := st.chat_input("What would you like to know?", max_chars=2000):
    st.session_state["messages"].append(HumanMessage(content=prompt))
    st.session_state["pending_prompt"] = prompt
    st.rerun()
